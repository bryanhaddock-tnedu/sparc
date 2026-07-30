from __future__ import annotations

import csv
import json
from dataclasses import dataclass
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from typing import Iterable
from uuid import uuid4
from zipfile import BadZipFile, ZipFile, ZIP_DEFLATED

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import func, select
from sqlalchemy.orm import Session, joinedload, selectinload

from app.models import (
    AppUser,
    Bucket,
    EstimationProfile,
    FiscalMonth,
    ForecastEntry,
    JiraProductMapping,
    JiraUserMapping,
    Product,
    ProductBudget,
    ProductJiraSpace,
    ProductTeamMember,
    RoadmapItem,
    RoadmapItemIssueLink,
    TeamMember,
)
from app.services.access_control import (
    UserRole,
    create_app_user,
    get_app_user_by_email,
    normalize_role,
    update_app_user,
)
from app.services.fiscal_year import ensure_fiscal_months
from app.services.product_org import product_org_pair_error
from app.services.roadmap import MANUAL_ROADMAP_LINK_SOURCE, _normalize_issue_key, _product_id_from_project_key
from app.services.slugs import unique_team_member_slug


@dataclass(frozen=True)
class AdminDataSet:
    key: str
    label: str
    sheet_name: str
    description: str
    default_selected: bool = True
    import_phase: str = "base"

    @property
    def csv_file_name(self) -> str:
        return f"{self.key}.csv"


EXPORT_DATASETS: tuple[AdminDataSet, ...] = (
    AdminDataSet("buckets", "Buckets", "Buckets", "Reference work buckets used by forecast and actual records."),
    AdminDataSet("estimation_profiles", "Estimation Profiles", "EstimationProfiles", "SPARC-owned estimation methods and policy rules."),
    AdminDataSet("team_members", "Team Members", "TeamMembers", "Roster, rates, employment type, and status."),
    AdminDataSet("products", "Products", "Products", "SPARC products and product-level details."),
    AdminDataSet("product_budgets", "Product Budgets", "ProductBudgets", "Fiscal-year-specific product budgets."),
    AdminDataSet("product_team_members", "Product Team Members", "ProductTeamMembers", "Manual product team assignments."),
    AdminDataSet("product_jira_spaces", "Product Jira Spaces", "ProductJiraSpaces", "Manual SPARC product-to-Jira project mappings."),
    AdminDataSet("jira_user_mappings", "Jira User Mappings", "JiraUserMappings", "Manual Jira user-to-roster mappings."),
    AdminDataSet("forecast_entries", "Forecast Entries", "ForecastEntries", "Manager-entered planning hours by product, person, bucket, and month."),
    AdminDataSet(
        "roadmap_item_overrides",
        "Roadmap Item Overrides",
        "RoadmapItemOverrides",
        "Manual SPARC Product and Bucket overrides. Import after Jira Roadmap sync.",
        import_phase="after_roadmap_sync",
    ),
    AdminDataSet(
        "roadmap_ticket_mappings",
        "Roadmap Ticket Mappings",
        "RoadmapTicketMappings",
        "Manual SPARC ticket-to-Roadmap Item associations. Import after Jira Roadmap sync.",
        import_phase="after_roadmap_sync",
    ),
    AdminDataSet(
        "user_access",
        "User Access Definitions",
        "UserAccess",
        "Optional users, roles, and Program Areas without passwords or Entra identity links.",
        default_selected=False,
    ),
)

EXCLUDED_PACKAGE_DATA = (
    "ActualEntry Jira worklogs",
    "SyncRun history",
    "JiraProjectCatalog refresh snapshots",
    "RoadmapItem Jira refresh rows",
    "EstimatedEntry generated estimates",
    "EstimationRun history",
    "EstimatedIssueAllocation audit rows",
    "AppUser password hashes, Entra identity links, and login history",
    "ForecastRecommendationDecision audit history",
    "AttributionChange correction audit history",
    "JiraProductMapping discovery cache",
)

_DATASET_BY_KEY = {dataset.key: dataset for dataset in EXPORT_DATASETS}
_SHEET_TO_DATASET_KEY = {dataset.sheet_name: dataset.key for dataset in EXPORT_DATASETS}
_CSV_TO_DATASET_KEY = {dataset.csv_file_name: dataset.key for dataset in EXPORT_DATASETS}


def export_options() -> list[dict[str, object]]:
    return [
        {
            "key": dataset.key,
            "label": dataset.label,
            "description": dataset.description,
            "default_selected": dataset.default_selected,
            "import_phase": dataset.import_phase,
        }
        for dataset in EXPORT_DATASETS
    ]


def normalize_dataset_keys(dataset_keys: Iterable[str] | None) -> list[str]:
    keys = [key for key in (dataset_keys or []) if key]
    if not keys:
        return [dataset.key for dataset in EXPORT_DATASETS if dataset.default_selected]
    unknown = sorted(set(keys) - set(_DATASET_BY_KEY))
    if unknown:
        raise ValueError(f"Unknown admin data set: {', '.join(unknown)}")
    return [dataset.key for dataset in EXPORT_DATASETS if dataset.key in keys]


def build_admin_data_export(db: Session, dataset_keys: Iterable[str] | None = None) -> BytesIO:
    selected_keys = normalize_dataset_keys(dataset_keys)
    workbook = Workbook()
    manifest = workbook.active
    manifest.title = "Manifest"
    _write_manifest(manifest, selected_keys)

    writers = {
        "buckets": _write_buckets,
        "estimation_profiles": _write_estimation_profiles,
        "team_members": _write_team_members,
        "products": _write_products,
        "product_budgets": _write_product_budgets,
        "product_team_members": _write_product_team_members,
        "product_jira_spaces": _write_product_jira_spaces,
        "jira_user_mappings": _write_jira_user_mappings,
        "forecast_entries": _write_forecast_entries,
        "roadmap_item_overrides": _write_roadmap_item_overrides,
        "roadmap_ticket_mappings": _write_roadmap_ticket_mappings,
        "user_access": _write_user_access,
    }

    for key in selected_keys:
        writers[key](workbook, db)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def build_admin_data_archive(db: Session, dataset_keys: Iterable[str] | None = None) -> BytesIO:
    selected_keys = normalize_dataset_keys(dataset_keys)
    workbook = load_workbook(build_admin_data_export(db, selected_keys), data_only=True)
    generated_at = datetime.now(timezone.utc).isoformat()
    package_id = str(uuid4())
    row_counts = {
        key: sum(
            1
            for values in workbook[_DATASET_BY_KEY[key].sheet_name].iter_rows(min_row=2, values_only=True)
            if any(value is not None and value != "" for value in values)
        )
        for key in selected_keys
    }

    output = BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
        manifest = {
            "export_type": "SPARC_ADMIN_DATA_PACKAGE",
            "version": 2,
            "format": "zip-csv",
            "package_id": package_id,
            "generated_at": generated_at,
            "included_datasets": selected_keys,
            "row_counts": row_counts,
            "excluded_package_data": list(EXCLUDED_PACKAGE_DATA),
            "files": {key: _DATASET_BY_KEY[key].csv_file_name for key in selected_keys},
        }
        archive.writestr("manifest.json", json.dumps(manifest, indent=2))

        for key in selected_keys:
            dataset = _DATASET_BY_KEY[key]
            sheet = workbook[dataset.sheet_name]
            csv_buffer = StringIO()
            writer = csv.writer(csv_buffer, lineterminator="\n")
            for values in sheet.iter_rows(values_only=True):
                writer.writerow(["" if value is None else value for value in values])
            archive.writestr(dataset.csv_file_name, csv_buffer.getvalue())

    output.seek(0)
    return output


def build_admin_data_json_package(db: Session, dataset_keys: Iterable[str] | None = None) -> dict[str, object]:
    selected_keys = normalize_dataset_keys(dataset_keys)
    workbook = load_workbook(build_admin_data_export(db, selected_keys), data_only=True)
    package_id = str(uuid4())
    data: dict[str, object] = {}
    row_counts: dict[str, int] = {}
    for key in selected_keys:
        dataset = _DATASET_BY_KEY[key]
        sheet = workbook[dataset.sheet_name]
        headers = [str(cell.value or "").strip() for cell in sheet[1]]
        rows = []
        for values in sheet.iter_rows(min_row=2, values_only=True):
            if all(value is None or value == "" for value in values):
                continue
            rows.append({headers[index]: value for index, value in enumerate(values) if index < len(headers)})
        row_counts[key] = len(rows)
        data[key] = {
            "label": dataset.label,
            "sheet_name": dataset.sheet_name,
            "headers": headers,
            "rows": rows,
        }
    return {
        "export_type": "SPARC_ADMIN_DATA_PACKAGE",
        "version": 2,
        "package_id": package_id,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "included_datasets": selected_keys,
        "row_counts": row_counts,
        "excluded_package_data": list(EXCLUDED_PACKAGE_DATA),
        "data": data,
    }


def import_admin_data_package(db: Session, content: bytes, dataset_keys: Iterable[str] | None = None) -> dict[str, object]:
    workbook = load_workbook(BytesIO(content), data_only=True)
    selected_keys = normalize_dataset_keys(dataset_keys)
    available_keys = [_SHEET_TO_DATASET_KEY[name] for name in workbook.sheetnames if name in _SHEET_TO_DATASET_KEY]
    keys_to_import = [key for key in selected_keys if key in available_keys]

    results = {key: {"key": key, "label": _DATASET_BY_KEY[key].label, "created": 0, "updated": 0, "skipped": 0, "failed": 0} for key in keys_to_import}
    errors: list[dict[str, object]] = []
    warnings: list[dict[str, object]] = []

    importers = {
        "buckets": _import_buckets,
        "estimation_profiles": _import_estimation_profiles,
        "team_members": _import_team_members,
        "products": _import_products,
        "product_budgets": _import_product_budgets,
        "product_team_members": _import_product_team_members,
        "product_jira_spaces": _import_product_jira_spaces,
        "jira_user_mappings": _import_jira_user_mappings,
        "forecast_entries": _import_forecast_entries,
        "roadmap_item_overrides": _import_roadmap_item_overrides,
        "roadmap_ticket_mappings": _import_roadmap_ticket_mappings,
        "user_access": _import_user_access,
    }

    # Dependencies matter: references first, then relationships, then forecast facts.
    for key in [dataset.key for dataset in EXPORT_DATASETS if dataset.key in keys_to_import]:
        importers[key](db, workbook, results[key], errors, warnings)
        db.flush()

    return {
        "datasets": list(results.values()),
        "errors": errors,
        "warnings": warnings,
        "excluded_jira_refresh_data": list(EXCLUDED_PACKAGE_DATA),
    }


def import_admin_data_archive(db: Session, content: bytes, dataset_keys: Iterable[str] | None = None) -> dict[str, object]:
    try:
        archive = ZipFile(BytesIO(content))
    except BadZipFile as exc:
        raise ValueError("Admin data package must be a SPARC zip export or .xlsx workbook") from exc

    with archive:
        manifest = _read_archive_manifest(archive)
        selected_keys = normalize_dataset_keys(dataset_keys) if dataset_keys is not None else _known_package_dataset_keys(manifest.get("included_datasets"))
        workbook = Workbook()
        workbook.active.title = "Manifest"

        for key in selected_keys:
            dataset = _DATASET_BY_KEY[key]
            if dataset.csv_file_name not in archive.namelist():
                continue
            sheet = workbook.create_sheet(dataset.sheet_name)
            text = archive.read(dataset.csv_file_name).decode("utf-8-sig")
            for row in csv.reader(StringIO(text)):
                sheet.append(row)

        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        result = import_admin_data_package(db, output.getvalue(), selected_keys)
        result["package"] = _package_info(manifest)
        return result


def import_admin_data_json_package(db: Session, content: bytes, dataset_keys: Iterable[str] | None = None) -> dict[str, object]:
    try:
        package = json.loads(content.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ValueError("JSON admin data package could not be parsed") from exc

    selected_keys = normalize_dataset_keys(dataset_keys) if dataset_keys is not None else _known_package_dataset_keys(package.get("included_datasets"))
    data = package.get("data")
    if not isinstance(data, dict):
        raise ValueError("JSON admin data package is missing data")

    workbook = Workbook()
    manifest = workbook.active
    manifest.title = "Manifest"
    _write_manifest(manifest, selected_keys)

    for key in selected_keys:
        dataset = _DATASET_BY_KEY[key]
        section = data.get(key)
        if not isinstance(section, dict):
            continue
        rows = section.get("rows", [])
        headers = section.get("headers", [])
        if not headers and isinstance(rows, list) and rows:
            headers = list(rows[0].keys())
        if not isinstance(headers, list):
            raise ValueError(f"JSON admin data package has invalid headers for {dataset.label}")
        if not isinstance(rows, list):
            raise ValueError(f"JSON admin data package has invalid rows for {dataset.label}")
        sheet = workbook.create_sheet(dataset.sheet_name)
        sheet.append(headers)
        for row in rows:
            if not isinstance(row, dict):
                continue
            sheet.append([row.get(header) for header in headers])
        _style_sheet(sheet)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    result = import_admin_data_package(db, output.getvalue(), selected_keys)
    result["package"] = _package_info(package)
    return result


def import_admin_data_content(db: Session, content: bytes, filename: str | None = None, dataset_keys: Iterable[str] | None = None) -> dict[str, object]:
    lower_name = (filename or "").lower()
    if lower_name.endswith(".zip"):
        return import_admin_data_archive(db, content, dataset_keys)
    if lower_name.endswith(".json") or content.lstrip().startswith(b"{"):
        return import_admin_data_json_package(db, content, dataset_keys)
    if not lower_name.endswith(".xlsx") and content.startswith(b"PK"):
        try:
            with ZipFile(BytesIO(content)) as archive:
                if "manifest.json" in archive.namelist() or any(name in _CSV_TO_DATASET_KEY for name in archive.namelist()):
                    return import_admin_data_archive(db, content, dataset_keys)
        except BadZipFile:
            pass
    return import_admin_data_package(db, content, dataset_keys)


def _read_archive_manifest(archive: ZipFile) -> dict[str, object]:
    if "manifest.json" not in archive.namelist():
        return {"included_datasets": [_CSV_TO_DATASET_KEY[name] for name in archive.namelist() if name in _CSV_TO_DATASET_KEY]}
    try:
        return json.loads(archive.read("manifest.json").decode("utf-8"))
    except json.JSONDecodeError as exc:
        raise ValueError("Admin data package manifest could not be parsed") from exc


def _known_package_dataset_keys(raw_keys: object) -> list[str]:
    keys = set(raw_keys) if isinstance(raw_keys, list) else set()
    return [dataset.key for dataset in EXPORT_DATASETS if dataset.key in keys]


def _package_info(manifest: dict[str, object]) -> dict[str, object]:
    raw_counts = manifest.get("row_counts")
    row_counts: dict[str, int] = {}
    if isinstance(raw_counts, dict):
        row_counts = {
            key: value
            for key, value in raw_counts.items()
            if isinstance(key, str) and isinstance(value, int)
        }
    return {
        "package_id": str(manifest.get("package_id") or "") or None,
        "version": int(manifest.get("version") or 1),
        "generated_at": str(manifest.get("generated_at") or "") or None,
        "row_counts": row_counts,
    }


def _write_manifest(sheet, selected_keys: list[str]) -> None:
    rows = [
        ("Export Type", "SPARC admin data package"),
        ("Generated At", datetime.now(timezone.utc).isoformat()),
        ("Included Data Sets", ", ".join(_DATASET_BY_KEY[key].label for key in selected_keys)),
        ("Excluded Package Data", ", ".join(EXCLUDED_PACKAGE_DATA)),
        ("Import Rule", "Rows are importable by natural keys like product name, team member name, Jira key, bucket code, fiscal year, and month sequence."),
        ("Note", "Use Jira Sync in SPARC to repopulate actual worklogs and Jira refresh data after importing this package."),
    ]
    _append_sheet_rows(sheet, ("Field", "Value"), rows)
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 140


def _write_buckets(workbook: Workbook, db: Session) -> None:
    rows = ((bucket.id, bucket.code, bucket.name) for bucket in db.scalars(select(Bucket).order_by(Bucket.id)).all())
    _append_sheet(workbook, "Buckets", ("source_bucket_id", "code", "name"), rows)


def _write_estimation_profiles(workbook: Workbook, db: Session) -> None:
    profiles = db.scalars(select(EstimationProfile).order_by(EstimationProfile.name)).all()
    rows = (
        (
            profile.id,
            profile.name,
            profile.description,
            profile.is_active,
            profile.method_version,
            _number(profile.monthly_capacity_hours),
            _number(profile.actual_completeness_threshold),
            profile.stale_ticket_window_days,
            profile.forecast_future_months,
            profile.future_month_average_window,
            profile.excluded_statuses,
            profile.low_activity_statuses,
            profile.excluded_jira_project_keys,
            profile.project_pause_dates,
            profile.work_type_field_priority,
            profile.story_point_weighting_enabled,
            profile.notes,
            _iso(profile.created_at),
            _iso(profile.updated_at),
        )
        for profile in profiles
    )
    _append_sheet(
        workbook,
        "EstimationProfiles",
        (
            "source_estimation_profile_id",
            "name",
            "description",
            "is_active",
            "method_version",
            "monthly_capacity_hours",
            "actual_completeness_threshold",
            "stale_ticket_window_days",
            "forecast_future_months",
            "future_month_average_window",
            "excluded_statuses",
            "low_activity_statuses",
            "excluded_jira_project_keys",
            "project_pause_dates",
            "work_type_field_priority",
            "story_point_weighting_enabled",
            "notes",
            "created_at",
            "updated_at",
        ),
        rows,
    )


def _write_team_members(workbook: Workbook, db: Session) -> None:
    members = db.scalars(select(TeamMember).order_by(TeamMember.name)).all()
    rows = (
        (
            member.id,
            member.staff_id,
            member.name,
            member.role,
            member.team,
            _number(member.bill_rate),
            member.employment_type,
            member.contracting_company,
            member.status,
            _iso(member.created_at),
            _iso(member.updated_at),
        )
        for member in members
    )
    _append_sheet(
        workbook,
        "TeamMembers",
        (
            "source_team_member_id",
            "staff_id",
            "name",
            "role",
            "team",
            "bill_rate",
            "employment_type",
            "contracting_company",
            "status",
            "created_at",
            "updated_at",
        ),
        rows,
    )


def _write_products(workbook: Workbook, db: Session) -> None:
    products = db.scalars(select(Product).order_by(Product.name)).all()
    rows = (
        (
            product.id,
            product.name,
            product.jira_space_key,
            product.description,
            product.office,
            product.division,
            _number(product.budget_amount),
            product.is_active,
            _iso(product.created_at),
            _iso(product.updated_at),
        )
        for product in products
    )
    _append_sheet(
        workbook,
        "Products",
        (
            "source_product_id",
            "name",
            "jira_space_key",
            "description",
            "office",
            "division",
            "legacy_budget_amount",
            "is_active",
            "created_at",
            "updated_at",
        ),
        rows,
    )


def _write_product_budgets(workbook: Workbook, db: Session) -> None:
    budgets = db.scalars(
        select(ProductBudget).options(joinedload(ProductBudget.product)).join(ProductBudget.product).order_by(Product.name, ProductBudget.fiscal_year)
    ).all()
    rows = (
        (
            budget.id,
            budget.product_id,
            budget.product.name,
            budget.fiscal_year,
            _number(budget.budget_amount),
            _iso(budget.created_at),
            _iso(budget.updated_at),
        )
        for budget in budgets
    )
    _append_sheet(
        workbook,
        "ProductBudgets",
        ("source_product_budget_id", "source_product_id", "product_name", "fiscal_year", "budget_amount", "created_at", "updated_at"),
        rows,
    )


def _write_product_team_members(workbook: Workbook, db: Session) -> None:
    assignments = db.scalars(
        select(ProductTeamMember)
        .options(
            joinedload(ProductTeamMember.product),
            joinedload(ProductTeamMember.team_member),
            joinedload(ProductTeamMember.default_bucket),
        )
        .join(ProductTeamMember.product)
        .join(ProductTeamMember.team_member)
        .order_by(Product.name, TeamMember.name)
    ).all()
    rows = (
        (
            assignment.id,
            assignment.product_id,
            assignment.product.name,
            assignment.team_member_id,
            assignment.team_member.staff_id,
            assignment.team_member.name,
            assignment.default_bucket.code if assignment.default_bucket else None,
            assignment.default_bucket.name if assignment.default_bucket else None,
            assignment.status,
            _iso(assignment.created_at),
            _iso(assignment.updated_at),
        )
        for assignment in assignments
    )
    _append_sheet(
        workbook,
        "ProductTeamMembers",
        (
            "source_assignment_id",
            "source_product_id",
            "product_name",
            "source_team_member_id",
            "team_member_staff_id",
            "team_member_name",
            "default_bucket_code",
            "default_bucket_name",
            "status",
            "created_at",
            "updated_at",
        ),
        rows,
    )


def _write_product_jira_spaces(workbook: Workbook, db: Session) -> None:
    spaces = db.scalars(
        select(ProductJiraSpace).options(joinedload(ProductJiraSpace.product)).join(ProductJiraSpace.product).order_by(Product.name, ProductJiraSpace.jira_project_key)
    ).all()
    rows = (
        (
            space.id,
            space.product_id,
            space.product.name,
            space.jira_project_catalog_id,
            space.jira_project_id,
            space.jira_project_key,
            space.jira_project_name,
            space.is_active,
            space.scope_jql,
            space.validation_status,
            space.validation_message,
            _iso(space.last_validated_at),
            _iso(space.created_at),
            _iso(space.updated_at),
        )
        for space in spaces
    )
    _append_sheet(
        workbook,
        "ProductJiraSpaces",
        (
            "source_product_jira_space_id",
            "source_product_id",
            "product_name",
            "jira_project_catalog_id",
            "jira_project_id",
            "jira_project_key",
            "jira_project_name",
            "is_active",
            "scope_jql",
            "validation_status",
            "validation_message",
            "last_validated_at",
            "created_at",
            "updated_at",
        ),
        rows,
    )


def _write_jira_user_mappings(workbook: Workbook, db: Session) -> None:
    mappings = db.scalars(
        select(JiraUserMapping).options(joinedload(JiraUserMapping.team_member)).order_by(JiraUserMapping.jira_display_name)
    ).all()
    rows = (
        (
            mapping.id,
            mapping.jira_account_id,
            mapping.jira_display_name,
            mapping.jira_email,
            mapping.team_member_id,
            mapping.team_member.staff_id if mapping.team_member else None,
            mapping.team_member.name if mapping.team_member else None,
            _iso(mapping.created_at),
            _iso(mapping.updated_at),
        )
        for mapping in mappings
    )
    _append_sheet(
        workbook,
        "JiraUserMappings",
        (
            "source_jira_user_mapping_id",
            "jira_account_id",
            "jira_display_name",
            "jira_email",
            "source_team_member_id",
            "team_member_staff_id",
            "team_member_name",
            "created_at",
            "updated_at",
        ),
        rows,
    )


def _write_jira_product_mappings(workbook: Workbook, db: Session) -> None:
    mappings = db.scalars(
        select(JiraProductMapping).options(joinedload(JiraProductMapping.product)).order_by(JiraProductMapping.jira_project_key)
    ).all()
    rows = (
        (
            mapping.id,
            mapping.jira_project_key,
            mapping.jira_project_name,
            mapping.product_id,
            mapping.product.name if mapping.product else None,
            _iso(mapping.created_at),
            _iso(mapping.updated_at),
        )
        for mapping in mappings
    )
    _append_sheet(
        workbook,
        "JiraProductMappings",
        (
            "source_jira_product_mapping_id",
            "jira_project_key",
            "jira_project_name",
            "source_product_id",
            "product_name",
            "created_at",
            "updated_at",
        ),
        rows,
    )


def _write_forecast_entries(workbook: Workbook, db: Session) -> None:
    forecasts = db.scalars(
        select(ForecastEntry)
        .options(
            joinedload(ForecastEntry.product),
            joinedload(ForecastEntry.team_member),
            joinedload(ForecastEntry.bucket),
            joinedload(ForecastEntry.fiscal_month),
        )
        .join(ForecastEntry.product)
        .join(ForecastEntry.team_member)
        .join(ForecastEntry.bucket)
        .join(ForecastEntry.fiscal_month)
        .order_by(FiscalMonth.fiscal_year, FiscalMonth.sequence, Product.name, TeamMember.name, Bucket.name)
    ).all()
    rows = (
        (
            forecast.id,
            forecast.product_id,
            forecast.product.name,
            forecast.team_member_id,
            forecast.team_member.staff_id,
            forecast.team_member.name,
            forecast.bucket_id,
            forecast.bucket.code,
            forecast.bucket.name,
            forecast.fiscal_month_id,
            forecast.fiscal_month.fiscal_year,
            forecast.fiscal_month.sequence,
            forecast.fiscal_month.label,
            forecast.fiscal_month.calendar_year,
            forecast.fiscal_month.calendar_month,
            _number(forecast.hours),
            _iso(forecast.created_at),
            _iso(forecast.updated_at),
        )
        for forecast in forecasts
    )
    _append_sheet(
        workbook,
        "ForecastEntries",
        (
            "source_forecast_entry_id",
            "source_product_id",
            "product_name",
            "source_team_member_id",
            "team_member_staff_id",
            "team_member_name",
            "source_bucket_id",
            "bucket_code",
            "bucket_name",
            "source_fiscal_month_id",
            "fiscal_year",
            "month_sequence",
            "month_label",
            "calendar_year",
            "calendar_month",
            "hours",
            "created_at",
            "updated_at",
        ),
        rows,
    )


def _write_roadmap_item_overrides(workbook: Workbook, db: Session) -> None:
    items = db.scalars(
        select(RoadmapItem)
        .options(joinedload(RoadmapItem.product), joinedload(RoadmapItem.bucket))
        .where(
            (RoadmapItem.product_mapping_source == "manual")
            | (RoadmapItem.bucket_mapping_source == "manual")
        )
        .order_by(RoadmapItem.fiscal_year, RoadmapItem.jira_issue_key)
    ).all()
    rows = (
        (
            item.source,
            item.fiscal_year,
            item.jira_issue_key,
            item.title,
            item.product_mapping_source == "manual",
            item.product.name if item.product else None,
            item.bucket_mapping_source == "manual",
            item.bucket.code if item.bucket else None,
            _iso(item.updated_at),
        )
        for item in items
    )
    _append_sheet(
        workbook,
        "RoadmapItemOverrides",
        (
            "roadmap_source",
            "fiscal_year",
            "roadmap_item_key",
            "roadmap_item_title",
            "has_product_override",
            "product_name",
            "has_bucket_override",
            "bucket_code",
            "updated_at",
        ),
        rows,
    )


def _write_roadmap_ticket_mappings(workbook: Workbook, db: Session) -> None:
    links = db.scalars(
        select(RoadmapItemIssueLink)
        .options(joinedload(RoadmapItemIssueLink.roadmap_item))
        .join(RoadmapItemIssueLink.roadmap_item)
        .where(RoadmapItemIssueLink.source == MANUAL_ROADMAP_LINK_SOURCE)
        .order_by(RoadmapItem.fiscal_year, RoadmapItemIssueLink.jira_issue_key)
    ).all()
    rows = (
        (
            link.roadmap_item.source,
            link.roadmap_item.fiscal_year,
            link.roadmap_item.jira_issue_key,
            link.roadmap_item.title,
            link.jira_issue_key,
            _iso(link.updated_at),
        )
        for link in links
    )
    _append_sheet(
        workbook,
        "RoadmapTicketMappings",
        (
            "roadmap_source",
            "fiscal_year",
            "roadmap_item_key",
            "roadmap_item_title",
            "ticket_key",
            "updated_at",
        ),
        rows,
    )


def _write_user_access(workbook: Workbook, db: Session) -> None:
    users = db.scalars(
        select(AppUser)
        .options(selectinload(AppUser.program_area_assignments))
        .order_by(func.lower(AppUser.email))
    ).all()
    rows = (
        (
            user.email,
            user.display_name,
            normalize_role(user.role).value,
            user.is_active,
            "; ".join(sorted(assignment.program_area for assignment in user.program_area_assignments)),
            _iso(user.created_at),
            _iso(user.updated_at),
        )
        for user in users
    )
    _append_sheet(
        workbook,
        "UserAccess",
        (
            "email",
            "display_name",
            "role",
            "is_active",
            "program_areas",
            "created_at",
            "updated_at",
        ),
        rows,
    )


def _import_buckets(
    db: Session,
    workbook,
    result: dict[str, int | str],
    errors: list[dict[str, object]],
    warnings: list[dict[str, object]],
) -> None:
    for row_number, row in _rows(workbook, "Buckets"):
        code = _text(row, "code")
        name = _text(row, "name")
        if not code or not name:
            _fail(result, errors, "Buckets", row_number, "Bucket code and name are required")
            continue
        bucket = db.scalar(select(Bucket).where(Bucket.code == code))
        if bucket is None:
            db.add(Bucket(code=code, name=name))
            result["created"] += 1
        else:
            bucket.name = name
            result["updated"] += 1


def _import_estimation_profiles(
    db: Session,
    workbook,
    result: dict[str, int | str],
    errors: list[dict[str, object]],
    warnings: list[dict[str, object]],
) -> None:
    for row_number, row in _rows(workbook, "EstimationProfiles"):
        name = _text(row, "name")
        method_version = _text(row, "method_version")
        if not name or not method_version:
            _fail(result, errors, "EstimationProfiles", row_number, "Profile name and method version are required")
            continue
        profile = db.scalar(select(EstimationProfile).where(EstimationProfile.name == name))
        created = profile is None
        if profile is None:
            profile = EstimationProfile(name=name, method_version=method_version)
            db.add(profile)
        profile.description = _text(row, "description")
        profile.is_active = _bool(row.get("is_active"), default=False)
        profile.method_version = method_version
        profile.monthly_capacity_hours = _decimal(row, "monthly_capacity_hours")
        profile.actual_completeness_threshold = _decimal(row, "actual_completeness_threshold")
        profile.stale_ticket_window_days = _int(row, "stale_ticket_window_days") or 10
        profile.forecast_future_months = _bool(row.get("forecast_future_months"), default=True)
        profile.future_month_average_window = _int(row, "future_month_average_window") or 2
        profile.excluded_statuses = _text(row, "excluded_statuses")
        profile.low_activity_statuses = _text(row, "low_activity_statuses")
        profile.excluded_jira_project_keys = _text(row, "excluded_jira_project_keys")
        profile.project_pause_dates = _text(row, "project_pause_dates")
        profile.work_type_field_priority = _text(row, "work_type_field_priority")
        profile.story_point_weighting_enabled = _bool(row.get("story_point_weighting_enabled"), default=True)
        profile.notes = _text(row, "notes")
        result["created" if created else "updated"] += 1


def _import_team_members(
    db: Session,
    workbook,
    result: dict[str, int | str],
    errors: list[dict[str, object]],
    warnings: list[dict[str, object]],
) -> None:
    for row_number, row in _rows(workbook, "TeamMembers"):
        name = _text(row, "name")
        if not name:
            _fail(result, errors, "TeamMembers", row_number, "Team member name is required")
            continue
        staff_id = _text(row, "staff_id")
        member = _find_team_member(db, row)
        created = member is None
        if member is None:
            member = TeamMember(name=name, role="", team="")
            db.add(member)
        member.staff_id = staff_id
        member.name = name
        member.role = _text(row, "role") or ""
        member.team = _text(row, "team") or ""
        member.bill_rate = _decimal(row, "bill_rate")
        member.employment_type = _text(row, "employment_type") or "Employee"
        member.contracting_company = _text(row, "contracting_company")
        member.status = _text(row, "status") or "active"
        if created or member.slug is None:
            member.slug = unique_team_member_slug(db, member.name, member.id)
        result["created" if created else "updated"] += 1


def _import_products(
    db: Session,
    workbook,
    result: dict[str, int | str],
    errors: list[dict[str, object]],
    warnings: list[dict[str, object]],
) -> None:
    for row_number, row in _rows(workbook, "Products"):
        name = _text(row, "name")
        if not name:
            _fail(result, errors, "Products", row_number, "Product name is required")
            continue
        office = _text(row, "office")
        division = _text(row, "division")
        if error := product_org_pair_error(office, division):
            _fail(result, errors, "Products", row_number, error)
            continue
        product = _find_product(db, row)
        created = product is None
        if product is None:
            product = Product(name=name)
            db.add(product)
        product.name = name
        product.jira_space_key = _text(row, "jira_space_key")
        product.description = _text(row, "description")
        product.office = office
        product.division = division
        product.budget_amount = _decimal(row, "legacy_budget_amount")
        product.is_active = _bool(row.get("is_active"), default=True)
        result["created" if created else "updated"] += 1


def _import_product_budgets(
    db: Session,
    workbook,
    result: dict[str, int | str],
    errors: list[dict[str, object]],
    warnings: list[dict[str, object]],
) -> None:
    for row_number, row in _rows(workbook, "ProductBudgets"):
        product = _find_product(db, row)
        fiscal_year = _int(row, "fiscal_year")
        if product is None or fiscal_year is None:
            _fail(result, errors, "ProductBudgets", row_number, "Product and fiscal year are required")
            continue
        budget = db.scalar(select(ProductBudget).where(ProductBudget.product_id == product.id, ProductBudget.fiscal_year == fiscal_year))
        created = budget is None
        if budget is None:
            budget = ProductBudget(product_id=product.id, fiscal_year=fiscal_year)
            db.add(budget)
        budget.budget_amount = _decimal(row, "budget_amount")
        result["created" if created else "updated"] += 1


def _import_product_team_members(
    db: Session,
    workbook,
    result: dict[str, int | str],
    errors: list[dict[str, object]],
    warnings: list[dict[str, object]],
) -> None:
    for row_number, row in _rows(workbook, "ProductTeamMembers"):
        product = _find_product(db, row)
        member = _find_team_member(db, row)
        if product is None or member is None:
            _fail(result, errors, "ProductTeamMembers", row_number, "Product and team member are required")
            continue
        assignment = db.scalar(select(ProductTeamMember).where(ProductTeamMember.product_id == product.id, ProductTeamMember.team_member_id == member.id))
        created = assignment is None
        if assignment is None:
            assignment = ProductTeamMember(product_id=product.id, team_member_id=member.id)
            db.add(assignment)
        bucket = _find_bucket(db, row, required=False)
        assignment.default_bucket_id = bucket.id if bucket else None
        assignment.status = _text(row, "status") or "active"
        result["created" if created else "updated"] += 1


def _import_product_jira_spaces(
    db: Session,
    workbook,
    result: dict[str, int | str],
    errors: list[dict[str, object]],
    warnings: list[dict[str, object]],
) -> None:
    for row_number, row in _rows(workbook, "ProductJiraSpaces"):
        product = _find_product(db, row)
        jira_project_key = _text(row, "jira_project_key")
        if product is None or not jira_project_key:
            _fail(result, errors, "ProductJiraSpaces", row_number, "Product and Jira project key are required")
            continue
        space = db.scalar(select(ProductJiraSpace).where(ProductJiraSpace.jira_project_key == jira_project_key))
        created = space is None
        if space is None:
            space = ProductJiraSpace(product_id=product.id, jira_project_key=jira_project_key)
            db.add(space)
        space.product_id = product.id
        space.jira_project_id = _text(row, "jira_project_id")
        space.jira_project_name = _text(row, "jira_project_name")
        space.is_active = _bool(row.get("is_active"), default=True)
        space.scope_jql = _text(row, "scope_jql")
        space.validation_status = _text(row, "validation_status") or "unknown"
        space.validation_message = _text(row, "validation_message")
        legacy_mapping = db.scalar(select(JiraProductMapping).where(JiraProductMapping.jira_project_key == jira_project_key))
        if legacy_mapping is None:
            legacy_mapping = JiraProductMapping(
                jira_project_key=jira_project_key,
                jira_project_name=space.jira_project_name or jira_project_key,
            )
            db.add(legacy_mapping)
        legacy_mapping.jira_project_name = space.jira_project_name or legacy_mapping.jira_project_name
        legacy_mapping.product_id = product.id if space.is_active else None
        result["created" if created else "updated"] += 1


def _import_jira_user_mappings(
    db: Session,
    workbook,
    result: dict[str, int | str],
    errors: list[dict[str, object]],
    warnings: list[dict[str, object]],
) -> None:
    for row_number, row in _rows(workbook, "JiraUserMappings"):
        jira_account_id = _text(row, "jira_account_id")
        if not jira_account_id:
            _fail(result, errors, "JiraUserMappings", row_number, "Jira account id is required")
            continue
        mapping = db.scalar(select(JiraUserMapping).where(JiraUserMapping.jira_account_id == jira_account_id))
        created = mapping is None
        if mapping is None:
            mapping = JiraUserMapping(jira_account_id=jira_account_id, jira_display_name="")
            db.add(mapping)
        member = _find_team_member(db, row, required=False)
        mapping.jira_display_name = _text(row, "jira_display_name") or jira_account_id
        mapping.jira_email = _text(row, "jira_email")
        mapping.team_member_id = member.id if member else None
        result["created" if created else "updated"] += 1


def _import_jira_product_mappings(
    db: Session,
    workbook,
    result: dict[str, int | str],
    errors: list[dict[str, object]],
    warnings: list[dict[str, object]],
) -> None:
    for row_number, row in _rows(workbook, "JiraProductMappings"):
        jira_project_key = _text(row, "jira_project_key")
        if not jira_project_key:
            _fail(result, errors, "JiraProductMappings", row_number, "Jira project key is required")
            continue
        mapping = db.scalar(select(JiraProductMapping).where(JiraProductMapping.jira_project_key == jira_project_key))
        created = mapping is None
        if mapping is None:
            mapping = JiraProductMapping(jira_project_key=jira_project_key, jira_project_name="")
            db.add(mapping)
        product = _find_product(db, row, required=False)
        mapping.jira_project_name = _text(row, "jira_project_name") or jira_project_key
        mapping.product_id = product.id if product else None
        result["created" if created else "updated"] += 1


def _import_forecast_entries(
    db: Session,
    workbook,
    result: dict[str, int | str],
    errors: list[dict[str, object]],
    warnings: list[dict[str, object]],
) -> None:
    rows = list(_rows(workbook, "ForecastEntries"))
    fiscal_years = sorted(
        {
            fiscal_year
            for _, row in rows
            if (fiscal_year := _int(row, "fiscal_year")) is not None
        }
    )
    for fiscal_year in fiscal_years:
        ensure_fiscal_months(db, fiscal_year)
    db.flush()

    products = db.scalars(select(Product)).all()
    members = db.scalars(select(TeamMember)).all()
    buckets = db.scalars(select(Bucket)).all()
    months = db.scalars(select(FiscalMonth).where(FiscalMonth.fiscal_year.in_(fiscal_years))).all() if fiscal_years else []
    assignments = db.scalars(select(ProductTeamMember)).all()
    entries = db.scalars(
        select(ForecastEntry)
        .join(ForecastEntry.fiscal_month)
        .where(FiscalMonth.fiscal_year.in_(fiscal_years))
    ).all() if fiscal_years else []

    products_by_name = {product.name: product for product in products}
    members_by_staff_id = {member.staff_id: member for member in members if member.staff_id}
    members_by_name = {member.name: member for member in members}
    buckets_by_code = {bucket.code: bucket for bucket in buckets}
    months_by_key = {(month.fiscal_year, month.sequence): month for month in months}
    assignments_by_key = {(assignment.product_id, assignment.team_member_id): assignment for assignment in assignments}
    entries_by_key = {
        (entry.product_id, entry.team_member_id, entry.bucket_id, entry.fiscal_month_id): entry
        for entry in entries
    }

    for row_number, row in rows:
        product = products_by_name.get(_text(row, "product_name") or "")
        staff_id = _text(row, "team_member_staff_id")
        member = members_by_staff_id.get(staff_id) if staff_id else None
        if member is None:
            member = members_by_name.get(_text(row, "team_member_name") or "")
        bucket = buckets_by_code.get(_text(row, "bucket_code") or "")
        fiscal_year = _int(row, "fiscal_year")
        month_sequence = _int(row, "month_sequence")
        month = months_by_key.get((fiscal_year, month_sequence)) if fiscal_year is not None and month_sequence is not None else None
        if product is None or member is None or bucket is None or month is None:
            missing = []
            if product is None:
                missing.append(f"Product '{_text(row, 'product_name') or ''}'")
            if member is None:
                missing.append(f"Team Member '{_text(row, 'team_member_name') or ''}'")
            if bucket is None:
                missing.append(f"Bucket '{_text(row, 'bucket_code') or ''}'")
            if month is None:
                missing.append(f"FY{fiscal_year or '?'} month {month_sequence or '?'}")
            _fail(result, errors, "ForecastEntries", row_number, f"Could not resolve {', '.join(missing)}")
            continue

        assignment_key = (product.id, member.id)
        if assignment_key not in assignments_by_key:
            assignment = ProductTeamMember(product_id=product.id, team_member_id=member.id, status="active")
            db.add(assignment)
            assignments_by_key[assignment_key] = assignment

        entry_key = (product.id, member.id, bucket.id, month.id)
        entry = entries_by_key.get(entry_key)
        if entry is None:
            entry = ForecastEntry(
                product_id=product.id,
                team_member_id=member.id,
                bucket_id=bucket.id,
                fiscal_month_id=month.id,
                hours=_decimal(row, "hours"),
            )
            db.add(entry)
            entries_by_key[entry_key] = entry
            result["created"] += 1
        else:
            entry.hours = _decimal(row, "hours")
            result["updated"] += 1


def _import_roadmap_item_overrides(
    db: Session,
    workbook,
    result: dict[str, int | str],
    errors: list[dict[str, object]],
    warnings: list[dict[str, object]],
) -> None:
    for row_number, row in _rows(workbook, "RoadmapItemOverrides"):
        item = _find_roadmap_item(db, row)
        if item is None:
            _warn(
                result,
                warnings,
                "RoadmapItemOverrides",
                row_number,
                "Roadmap Item not found. Run Jira Roadmap sync, then import this data set again.",
            )
            continue

        changed = False
        if _bool(row.get("has_product_override")):
            product_name = _text(row, "product_name")
            product = db.scalar(select(Product).where(Product.name == product_name)) if product_name else None
            if product is None:
                _fail(result, errors, "RoadmapItemOverrides", row_number, "Product override could not be resolved")
                continue
            item.product_id = product.id
            item.product_mapping_source = "manual"
            changed = True

        if _bool(row.get("has_bucket_override")):
            bucket_code = _text(row, "bucket_code")
            bucket = db.scalar(select(Bucket).where(Bucket.code == bucket_code)) if bucket_code else None
            if bucket is None:
                _fail(result, errors, "RoadmapItemOverrides", row_number, "Bucket override could not be resolved")
                continue
            item.bucket_id = bucket.id
            item.bucket_mapping_source = "manual"
            changed = True

        if changed:
            result["updated"] += 1
        else:
            result["skipped"] += 1


def _import_roadmap_ticket_mappings(
    db: Session,
    workbook,
    result: dict[str, int | str],
    errors: list[dict[str, object]],
    warnings: list[dict[str, object]],
) -> None:
    for row_number, row in _rows(workbook, "RoadmapTicketMappings"):
        item = _find_roadmap_item(db, row)
        if item is None:
            _warn(
                result,
                warnings,
                "RoadmapTicketMappings",
                row_number,
                "Roadmap Item not found. Run Jira Roadmap sync, then import this data set again.",
            )
            continue
        ticket_key = _normalize_issue_key(_text(row, "ticket_key"))
        if not ticket_key:
            _fail(result, errors, "RoadmapTicketMappings", row_number, "Ticket key is required")
            continue

        existing_links = db.scalars(
            select(RoadmapItemIssueLink)
            .options(joinedload(RoadmapItemIssueLink.roadmap_item))
            .where(RoadmapItemIssueLink.jira_issue_key == ticket_key)
        ).all()
        current_year_links = [link for link in existing_links if link.roadmap_item.fiscal_year == item.fiscal_year]
        project_key = ticket_key.split("-", 1)[0]
        product_id = _product_id_from_project_key(db, project_key)

        if len(current_year_links) == 1 and current_year_links[0].roadmap_item_id == item.id:
            link = current_year_links[0]
            link.source = MANUAL_ROADMAP_LINK_SOURCE
            link.product_id = product_id
            link.jira_project_key = project_key
            result["updated"] += 1
            continue

        for link in current_year_links:
            db.delete(link)
        if current_year_links:
            db.flush()
        db.add(
            RoadmapItemIssueLink(
                roadmap_item_id=item.id,
                product_id=product_id,
                jira_issue_key=ticket_key,
                jira_project_key=project_key,
                source=MANUAL_ROADMAP_LINK_SOURCE,
                last_synced_at=datetime.now(timezone.utc),
            )
        )
        result["updated" if current_year_links else "created"] += 1


def _import_user_access(
    db: Session,
    workbook,
    result: dict[str, int | str],
    errors: list[dict[str, object]],
    warnings: list[dict[str, object]],
) -> None:
    for row_number, row in _rows(workbook, "UserAccess"):
        email = _text(row, "email")
        display_name = _text(row, "display_name")
        role_value = _text(row, "role")
        if not email or not display_name or not role_value:
            _fail(result, errors, "UserAccess", row_number, "Email, display name, and role are required")
            continue
        try:
            role = normalize_role(role_value)
        except ValueError as exc:
            _fail(result, errors, "UserAccess", row_number, str(exc))
            continue

        program_areas = [
            value.strip()
            for value in (_text(row, "program_areas") or "").split(";")
            if value.strip()
        ]
        active = _bool(row.get("is_active"), default=True)
        existing = get_app_user_by_email(db, email)
        if (
            existing is not None
            and existing.is_active
            and normalize_role(existing.role) == UserRole.ADMIN
            and (not active or role != UserRole.ADMIN)
            and _active_admin_count(db) <= 1
        ):
            _fail(result, errors, "UserAccess", row_number, "Cannot disable or demote the final active SPARC Admin")
            continue

        try:
            if existing is None:
                create_app_user(
                    db,
                    email=email,
                    display_name=display_name,
                    role=role.value,
                    program_areas=program_areas,
                    active=active,
                    local_login_enabled=False,
                )
                result["created"] += 1
                _warn(
                    result,
                    warnings,
                    "UserAccess",
                    row_number,
                    "User created with local login disabled and no Entra identity link.",
                    increment_skipped=False,
                )
            else:
                update_app_user(
                    db,
                    existing,
                    display_name=display_name,
                    role=role.value,
                    program_areas=program_areas,
                    active=active,
                )
                result["updated"] += 1
        except ValueError as exc:
            _fail(result, errors, "UserAccess", row_number, str(exc))


def _append_sheet(workbook: Workbook, sheet_name: str, headers: tuple[str, ...], rows: Iterable[tuple[object, ...]]) -> None:
    sheet = workbook.create_sheet(sheet_name)
    _append_sheet_rows(sheet, headers, rows)


def _append_sheet_rows(sheet, headers: tuple[str, ...], rows: Iterable[tuple[object, ...]]) -> None:
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    _style_sheet(sheet)


def _style_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="002D72")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="left")
    sheet.freeze_panes = "A2"
    if sheet.max_row >= 1 and sheet.max_column >= 1:
        sheet.auto_filter.ref = sheet.dimensions
    for column_cells in sheet.columns:
        values = [str(cell.value) for cell in column_cells if cell.value is not None]
        width = max([len(value) for value in values] + [12])
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max(width + 2, 12), 48)


def _rows(workbook, sheet_name: str) -> Iterable[tuple[int, dict[str, object]]]:
    if sheet_name not in workbook.sheetnames:
        return []
    sheet = workbook[sheet_name]
    headers = [str(cell.value or "").strip() for cell in sheet[1]]
    row_dicts = []
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if all(value is None or value == "" for value in values):
            continue
        row_dicts.append((row_number, {headers[index]: value for index, value in enumerate(values) if index < len(headers)}))
    return row_dicts


def _find_product(db: Session, row: dict[str, object], required: bool = True) -> Product | None:
    name = _text(row, "product_name") or _text(row, "name")
    if not name:
        return None
    return db.scalar(select(Product).where(Product.name == name))


def _find_team_member(db: Session, row: dict[str, object], required: bool = True) -> TeamMember | None:
    staff_id = _text(row, "team_member_staff_id") or _text(row, "staff_id")
    if staff_id:
        member = db.scalar(select(TeamMember).where(TeamMember.staff_id == staff_id))
        if member is not None:
            return member
    name = _text(row, "team_member_name") or _text(row, "name")
    if not name:
        return None
    return db.scalar(select(TeamMember).where(TeamMember.name == name))


def _find_bucket(db: Session, row: dict[str, object], required: bool = True) -> Bucket | None:
    code = _text(row, "bucket_code") or _text(row, "code") or _text(row, "default_bucket_code")
    if code:
        bucket = db.scalar(select(Bucket).where(Bucket.code == code))
        if bucket is not None:
            return bucket
    name = _text(row, "bucket_name") or _text(row, "name") or _text(row, "default_bucket_name")
    if not name:
        return None
    return db.scalar(select(Bucket).where(Bucket.name == name))


def _find_roadmap_item(db: Session, row: dict[str, object]) -> RoadmapItem | None:
    source = _text(row, "roadmap_source") or "jira_product_discovery"
    fiscal_year = _int(row, "fiscal_year")
    issue_key = _normalize_issue_key(_text(row, "roadmap_item_key"))
    if fiscal_year is None or not issue_key:
        return None
    return db.scalar(
        select(RoadmapItem).where(
            RoadmapItem.source == source,
            RoadmapItem.fiscal_year == fiscal_year,
            RoadmapItem.jira_issue_key == issue_key,
        )
    )


def _active_admin_count(db: Session) -> int:
    return int(
        db.scalar(
            select(func.count())
            .select_from(AppUser)
            .where(AppUser.is_active.is_(True), AppUser.role == UserRole.ADMIN.value)
        )
        or 0
    )


def _fail(result: dict[str, int | str], errors: list[dict[str, object]], sheet: str, row: int, message: str) -> None:
    result["failed"] += 1
    errors.append({"sheet": sheet, "row": row, "message": message})


def _warn(
    result: dict[str, int | str],
    warnings: list[dict[str, object]],
    sheet: str,
    row: int,
    message: str,
    *,
    increment_skipped: bool = True,
) -> None:
    if increment_skipped:
        result["skipped"] += 1
    warnings.append({"sheet": sheet, "row": row, "message": message})


def _text(row: dict[str, object], key: str) -> str | None:
    value = row.get(key)
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _int(row: dict[str, object], key: str) -> int | None:
    value = row.get(key)
    if value is None or value == "":
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _decimal(row: dict[str, object], key: str) -> Decimal:
    value = row.get(key)
    if value is None or value == "":
        return Decimal("0.00")
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError):
        return Decimal("0.00")


def _bool(value: object, default: bool = False) -> bool:
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in {"1", "true", "yes", "y", "active"}


def _number(value: Decimal | int | float | None) -> float:
    return float(value or 0)


def _iso(value: object) -> str | None:
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)
