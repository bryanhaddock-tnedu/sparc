import csv
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from typing import Any

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models import TeamMember
from app.services.team_members import create_team_member, find_existing_member, normalize_name, update_team_member

REQUIRED_COLUMNS = {"First Name", "Role", "Team", "Employment Type", "Contracting Company"}
COLUMN_ALIASES = {
    "bill rate": "Bill Rate",
    "company": "Contracting Company",
    "contracting company": "Contracting Company",
    "employment type": "Employment Type",
    "employee type": "Employment Type",
    "first name": "First Name",
    "firstname": "First Name",
    "given name": "First Name",
    "last name": "Last Name",
    "lastname": "Last Name",
    "role": "Role",
    "staff id": "Staff ID",
    "team": "Team",
    "team member id": "Staff ID",
}
FTE_EMPLOYMENT_TYPES = {"employee", "fte", "full time", "full-time", "fulltime"}


@dataclass(frozen=True)
class ParsedRosterRow:
    row_number: int
    values: dict[str, Any]


def import_team_members(db: Session, *, filename: str, content: bytes) -> dict[str, Any]:
    rows = _read_rows(filename, content)
    result: dict[str, Any] = {"created": 0, "updated": 0, "skipped": 0, "failed": 0, "rows": [], "errors": []}

    for row in rows:
        try:
            payload = _row_to_payload(row.values)
        except ValueError as exc:
            result["failed"] += 1
            result["errors"].append({"row": row.row_number, "message": str(exc)})
            continue

        existing = find_existing_member(db, staff_id=payload.get("staff_id"), name=str(payload["name"]))
        if existing is None:
            member = create_team_member(db, payload)
            result["created"] += 1
            action = "created"
        else:
            update_team_member(db, existing, payload)
            member = existing
            result["updated"] += 1
            action = "updated"
        result["rows"].append(_result_row(row.row_number, member, action))

    db.flush()
    return result


def _read_rows(filename: str, content: bytes) -> list[ParsedRosterRow]:
    lowered = filename.lower()
    if lowered.endswith(".csv"):
        decoded = content.decode("utf-8-sig")
        return _parse_tabular_rows(list(csv.reader(StringIO(decoded))))
    if lowered.endswith(".xlsx"):
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        return _parse_tabular_rows(list(sheet.iter_rows(values_only=True)))
    raise ValueError("Upload must be a .csv or .xlsx file")


def _parse_tabular_rows(raw_rows: list[tuple[Any, ...] | list[Any]]) -> list[ParsedRosterRow]:
    header_index, header_map = _find_header(raw_rows)
    rows: list[ParsedRosterRow] = []
    for offset, values in enumerate(raw_rows[header_index + 1 :], start=header_index + 2):
        if _is_blank_row(values):
            continue
        row = {column: values[index] if index < len(values) else None for column, index in header_map.items()}
        rows.append(ParsedRosterRow(row_number=offset, values=row))
    return rows


def _find_header(raw_rows: list[tuple[Any, ...] | list[Any]]) -> tuple[int, dict[str, int]]:
    best_missing = sorted(REQUIRED_COLUMNS)
    for index, values in enumerate(raw_rows):
        header_map = _canonical_header_map(values)
        missing = sorted(REQUIRED_COLUMNS - set(header_map))
        if not missing:
            return index, header_map
        if len(missing) < len(best_missing):
            best_missing = missing
    raise ValueError(f"Missing required columns: {', '.join(best_missing)}")


def _canonical_header_map(values: tuple[Any, ...] | list[Any]) -> dict[str, int]:
    header_map: dict[str, int] = {}
    for index, value in enumerate(values):
        canonical = _canonical_column(value)
        if canonical and canonical not in header_map:
            header_map[canonical] = index
    return header_map


def _canonical_column(value: Any) -> str | None:
    text = _optional_text(value)
    if text is None:
        return None
    return COLUMN_ALIASES.get(_normalize_header(text))


def _normalize_header(value: str) -> str:
    return " ".join(value.replace("_", " ").strip().lower().split())


def _is_blank_row(values: tuple[Any, ...] | list[Any]) -> bool:
    return all(_optional_text(value) is None for value in values)


def _row_to_payload(row: dict[str, Any]) -> dict[str, Any]:
    first_name = _required_text(row, "First Name")
    last_name = _optional_text(row.get("Last Name"))
    role = _required_text(row, "Role")
    team = _required_text(row, "Team")
    employment_type = _required_text(row, "Employment Type")
    contracting_company = _optional_text(row.get("Contracting Company"))
    staff_id = _optional_text(row.get("Staff ID"))
    bill_rate = _parse_bill_rate(row.get("Bill Rate"), employment_type)

    name = normalize_name(" ".join(part for part in [first_name, last_name] if part))

    return {
        "staff_id": staff_id,
        "name": name,
        "role": role,
        "team": team,
        "bill_rate": bill_rate,
        "employment_type": employment_type,
        "contracting_company": contracting_company,
        "status": "active",
    }


def _parse_bill_rate(value: Any, employment_type: str) -> Decimal:
    text = _optional_text(value)
    if text is None:
        if _is_fte(employment_type):
            return Decimal("0.00")
        raise ValueError("Bill Rate is required for contractors")
    try:
        bill_rate = Decimal(text.replace("$", "").replace(",", "").strip())
    except (InvalidOperation, AttributeError) as exc:
        raise ValueError("Bill Rate must be a valid number") from exc
    if bill_rate < 0:
        raise ValueError("Bill Rate cannot be negative")
    return bill_rate


def _is_fte(employment_type: str) -> bool:
    return _normalize_header(employment_type) in FTE_EMPLOYMENT_TYPES


def _required_text(row: dict[str, Any], column: str) -> str:
    value = _optional_text(row.get(column))
    if not value:
        raise ValueError(f"{column} is required")
    return value


def _optional_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _result_row(row: int, member: TeamMember, action: str) -> dict[str, Any]:
    return {
        "row": row,
        "team_member_id": member.id,
        "staff_id": member.staff_id,
        "name": member.name,
        "action": action,
    }
